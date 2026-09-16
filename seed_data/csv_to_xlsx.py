"""
csv_to_xlsx.py
把 seed_data/ 目录下的 CSV 转成 .xlsx，方便老师在小程序里上传导入。

表头约定（与 publishAssignment 云函数读取顺序一致）：
    第 1 列：题目
    第 2 列：参考答案
    第 3 列：截止时间（支持 2026-12-31 或 2026-12-31 23:59）

用法：
    python3.11 miniapp/seed_data/csv_to_xlsx.py
"""
import csv
import sys
from pathlib import Path

try:
    from openpyxl import Workbook
    from openpyxl.styles import Font, PatternFill, Alignment
except ImportError:
    print("❌ 缺 openpyxl，先装：python3.11 -m pip install openpyxl")
    sys.exit(1)

HERE = Path(__file__).parent

HEADER_FONT = Font(bold=True, size=12)
HEADER_FILL = PatternFill(start_color="E0E0E0", end_color="E0E0E0", fill_type="solid")
WRAP = Alignment(wrap_text=True, vertical="top")

# 列宽：题目短一些、参考答案长一些、截止时间固定
COL_WIDTHS = {
    "题目": 32,
    "参考答案": 60,
    "截止时间": 18,
}

def convert(csv_path: Path) -> Path:
    wb = Workbook()
    ws = wb.active
    ws.title = "背诵题目"

    with csv_path.open("r", encoding="utf-8-sig", newline="") as f:
        reader = csv.reader(f)
        rows = list(reader)

    if not rows:
        print(f"⚠️ {csv_path.name} 是空的，跳过")
        return None

    header = rows[0]
    data_rows = rows[1:]

    # 表头
    for col_idx, name in enumerate(header, start=1):
        cell = ws.cell(row=1, column=col_idx, value=name)
        cell.font = HEADER_FONT
        cell.fill = HEADER_FILL
        cell.alignment = WRAP
        if name in COL_WIDTHS:
            ws.column_dimensions[cell.column_letter].width = COL_WIDTHS[name]

    # 数据（多空行/分点的参考答案在单元格里自动换行展示）
    for r_idx, row in enumerate(data_rows, start=2):
        for c_idx, value in enumerate(row, start=1):
            cell = ws.cell(row=r_idx, column=c_idx, value=value)
            cell.alignment = WRAP

    ws.freeze_panes = "A2"  # 冻结首行

    xlsx_path = csv_path.with_suffix(".xlsx")
    wb.save(xlsx_path)
    print(f"✅ {csv_path.name} -> {xlsx_path.name}（{len(data_rows)} 题）")
    return xlsx_path

if __name__ == "__main__":
    csv_files = sorted(HERE.glob("*.csv"))
    if not csv_files:
        print(f"⚠️ {HERE} 下没找到 .csv 文件")
        sys.exit(0)

    for csv_path in csv_files:
        convert(csv_path)

    print("\n🎉 全部转换完成，xlsx 就在 CSV 旁边，可以上传到小程序了。")
